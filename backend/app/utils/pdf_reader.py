import os
import pdfplumber
from docx import Document


def extract_text(path: str) -> str:
    """
    Extract plain text from a CV file.
    Supports:
        - PDF  (.pdf)
        - Word (.docx)
    Raises ValueError for unsupported formats.
    """
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return _extract_from_pdf(path)

    if ext == ".docx":
        return _extract_from_docx(path)

    raise ValueError(
        f"Unsupported file type '{ext}'. "
        "Please upload a PDF or DOCX file."
    )


# -----------------------------------------
# PDF extraction
# -----------------------------------------
def _extract_from_pdf(path: str) -> str:
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() or ""
    return text


# -----------------------------------------
# DOCX extraction  (paragraphs + tables)
# -----------------------------------------
def _extract_from_docx(path: str) -> str:
    doc = Document(path)
    parts = []

    # 1 - Normal paragraphs
    for para in doc.paragraphs:
        line = para.text.strip()
        if line:
            parts.append(line)

    # 2 - Table cells (skills/education often live in tables)
    for table in doc.tables:
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                parts.append(" | ".join(row_texts))

    return "\n".join(parts)